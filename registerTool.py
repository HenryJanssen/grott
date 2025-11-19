"""
Register Tool - Export/Import register tables to/from JSON and XLSX formats
Supports datalogger, inverter time/date, and inverter configuration registers
"""

import json
import os
from pathlib import Path

# Try to import openpyxl for XLSX support
try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font
    XLSX_AVAILABLE = True
except ImportError:
    XLSX_AVAILABLE = False


# Define register tables with readOnStart column
datalogger_registers_table = (
    {"register": 4, "name": "Interval", "description": "update interval, Ascii, e.g 5 or 1 or 0.5", "readOnStart": False},
    {"register": 17, "name": "growatt_ip", "description": "Growatt server ip addres, Ascii, set for redirection to Grott e.g. 192.168.0.206", "readOnStart": False},
    {"register": 18, "name": "growatt_port", "description": "Growatt server Port, Num, set for redirection to Grott e.g. 5279", "readOnStart": False},
    {"register": 31, "name": "datetime", "description": "current date-time, Ascii, e.g 2022-05-17 21:01:50", "readOnStart": False}
)

inverter_timedate_registers_table = (
    {"register": 45, "name": "Year", "description": "Inverter time: 4 digit year when read, offset from 2000 when written", "value": "yyyy", "unit": "int", "initial": "", "readOnStart": False},
    {"register": 46, "name": "Month", "description": "Inverter time: month", "value": "1-12", "unit": "int", "initial": "", "readOnStart": False},
    {"register": 47, "name": "Day", "description": "Inverter time: day", "value": "1-31", "unit": "int", "initial": "", "readOnStart": False},
    {"register": 48, "name": "Hour", "description": "Inverter time: hour", "value": "0-23", "unit": "int", "initial": "", "readOnStart": False},
    {"register": 49, "name": "Minute", "description": "Inverter time: minute", "value": "0-59", "unit": "int", "initial": "", "readOnStart": False},
    {"register": 50, "name": "Second", "description": "Inverter time: second", "value": "0-59", "unit": "int", "initial": "", "readOnStart": False}
)

inverter_registers_table = (
    {"register": 1000, "name": "Float charge current limit", "description": "When charge current battery need is lower than this value, enter nto float charge", "value": "", "unit": "0.1", "initial": "600", "readOnStart": False},
    {"register": 1044, "name": "Priority", "description": "ForceChrEn / ForceDischrEn Load first / Bat first / Grid first", "value": "0:Load (default) 1:Battery 2:Grid", "unit": "int", "initial": "0", "readOnStart": False},
    {"register": 1060, "name": "BuckUpsFunEn", "description": "Ups function enable or disable", "value": "Enable: 1 Disable: 0", "unit": "int", "initial": "", "readOnStart": False},
    {"register": 1061, "name": "BuckUPSVoltSet", "description": "UPS output voltage", "value": "0:230 1:208 2:240", "unit": "int", "initial": "230v", "readOnStart": False},
    {"register": 1062, "name": "UPSFreqSet", "description": "UPS output frequency", "value": "0:50Hz 1:60Hz", "unit": "int", "initial": "50Hz", "readOnStart": False},
    {"register": 1070, "name": "GridFirstDischargePowerRate", "description": "Discharge Power Rate when Grid First", "value": "0-100", "unit": "1%", "initial": "", "readOnStart": False},
    {"register": 1071, "name": "GridFirstStopSOC", "description": "Stop Discharge soc when Grid First", "value": "0-100", "unit": "1%", "initial": "", "readOnStart": False},
    {"register": 1080, "name": "Grid First Start Time 1", "description": "High eight bit: hour Low eight bit: minute", "value": "0-23 0-59", "unit": "hextime", "initial": "", "readOnStart": False},
    {"register": 1081, "name": "Grid First Stop Time 1", "description": "High eight bit: hour Low eight bit: minute", "value": "0-23 0-59", "unit": "hextime", "initial": "", "readOnStart": False},
    {"register": 1082, "name": "Grid First Stop Switch 1", "description": "Enable: 1 Disable: 0", "value": "0 or 1", "unit": "int", "initial": "", "readOnStart": False},
    {"register": 1083, "name": "Grid First Start Time 2", "description": "High eight bit: hour Low eight bit: minute", "value": "0-23 0-59", "unit": "hextime", "initial": "", "readOnStart": False},
    {"register": 1084, "name": "Grid First Stop Time 2", "description": "High eight bit: hour Low eight bit: minute", "value": "0-23 0-59", "unit": "hextime", "initial": "", "readOnStart": False},
    {"register": 1085, "name": "Grid First Stop Switch 2", "description": "Enable: 1 Disable: 0 ForceDischarge Switch&LCD_SET_FORCE_TRUE_2)==LCD_SET_FORCE_TRUE_2", "value": "0 or 1", "unit": "int", "initial": "", "readOnStart": False},
    {"register": 1086, "name": "Grid First Start Time 3", "description": "High eight bit: hour Low eight bit: minute", "value": "0-23 0-59", "unit": "hextime", "initial": "", "readOnStart": False},
    {"register": 1087, "name": "Grid First Stop Time 3", "description": "High eight bit: hour Low eight bit: minute", "value": "0-23 0-59", "unit": "hextime", "initial": "", "readOnStart": False},
    {"register": 1088, "name": "Grid First Stop Switch 3", "description": "Enable: 1 Disable: 0", "value": "0 or 1", "unit": "int", "initial": "", "readOnStart": False},
    {"register": 1090, "name": "Bat FirstPower Rate", "description": "Charge Power Rate when Bat First", "value": "0-100", "unit": "1%", "initial": "", "readOnStart": False},
    {"register": 1091, "name": "Bat First stop SOC", "description": "Stop Charge soc when Bat First", "value": "0-100", "unit": "1%", "initial": "", "readOnStart": False},
    {"register": 1092, "name": "AC charge Switch", "description": "When Bat First Enable: 1 Disable: 0", "value": "0 or 1", "unit": "int", "initial": "", "readOnStart": False},
    {"register": 1100, "name": "Bat First Start Time 1", "description": "High eight bit: hour Low eight bit: minute", "value": "0-23 0-59", "unit": "hextime", "initial": "", "readOnStart": False},
    {"register": 1101, "name": "Bat First Stop Time 1", "description": "High eight bit: hour Low eight bit: minute", "value": "0-23 0-59", "unit": "hextime", "initial": "", "readOnStart": False},
    {"register": 1102, "name": "Bat First on/off Switch 1", "description": "Enable: 1 Disable: 0", "value": "0 or 1", "unit": "int", "initial": "", "readOnStart": False},
    {"register": 1103, "name": "Bat First Start Time 2", "description": "High eight bit: hour Low eight bit: minute", "value": "0-23 0-59", "unit": "hextime", "initial": "", "readOnStart": False},
    {"register": 1104, "name": "Bat First Stop Time 2", "description": "High eight bit: hour Low eight bit: minute", "value": "0-23 0-59", "unit": "hextime", "initial": "", "readOnStart": False},
    {"register": 1105, "name": "Bat First on/off Switch 2", "description": "Enable: 1 Disable: 0", "value": "0 or 1", "unit": "int", "initial": "", "readOnStart": False},
    {"register": 1106, "name": "Bat First Start Time 3", "description": "High eight bit: hour Low eight bit: minute", "value": "0-23 0-59", "unit": "hextime", "initial": "", "readOnStart": False},
    {"register": 1107, "name": "Bat First Stop Time 3", "description": "High eight bit: hour Low eight bit: minute", "value": "0-23 0-59", "unit": "hextime", "initial": "", "readOnStart": False},
    {"register": 1108, "name": "Bat First on/off Switch 3", "description": "Enable: 1 Disable: 0", "value": "0 or 1", "unit": "int", "initial": "", "readOnStart": False},
    {"register": 1109, "name": "NoName", "description": "Load First Discharge Stopped Soc", "value": "0-100", "unit": "int", "initial": "", "readOnStart": False},
    {"register": 1110, "name": "Load First Start Time 1", "description": "High eight bit: hour Low eight bit: minute", "value": "0-23 0-59", "unit": "hextime", "initial": "", "readOnStart": False},
    {"register": 1111, "name": "Load First Stop Time 1", "description": "High eight bit: hour Low eight bit: minute", "value": "0-23 0-59", "unit": "hextime", "initial": "", "readOnStart": False},
    {"register": 1112, "name": "Load First on/off Switch 1", "description": "Enable: 1 Disable: 0", "value": "0 or 1", "unit": "int", "initial": "", "readOnStart": False},
    {"register": 1113, "name": "Load First Start Time 2", "description": "High eight bit: hour Low eight bit: minute", "value": "0-23 0-59", "unit": "hextime", "initial": "", "readOnStart": False},
    {"register": 1114, "name": "Load First Stop Time 2", "description": "High eight bit: hour Low eight bit: minute", "value": "0-23 0-59", "unit": "hextime", "initial": "", "readOnStart": False},
    {"register": 1115, "name": "Load First on/off Switch 2", "description": "Enable: 1 Disable: 0", "value": "0 or 1", "unit": "int", "initial": "", "readOnStart": False},
    {"register": 1116, "name": "Load First Start Time 3", "description": "High eight bit: hour Low eight bit: minute", "value": "0-23 0-59", "unit": "hextime", "initial": "", "readOnStart": False},
    {"register": 1117, "name": "Load First Stop Time 3", "description": "High eight bit: hour Low eight bit: minute", "value": "0-23 0-59", "unit": "hextime", "initial": "", "readOnStart": False},
    {"register": 1118, "name": "Load First on/off Switch 3", "description": "Enable: 1 Disable: 0", "value": "0 or 1", "unit": "int", "initial": "", "readOnStart": False}
)

all_registers_table = (
    datalogger_registers_table,
    inverter_timedate_registers_table,
    inverter_registers_table
)


class RegisterExporter:
    """Export register tables to JSON and XLSX formats"""
    
    @staticmethod
    def export_to_json(registers, filename):
        """Export registers to JSON file"""
        data = list(registers)
        with open(filename, 'w', encoding='utf-8') as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
        print(f"✓ Exported {len(data)} registers to {filename}")
        return filename
    
    @staticmethod
    def export_to_xlsx(registers, filename):
        """Export registers to XLSX file"""
        if not XLSX_AVAILABLE:
            raise ImportError("openpyxl is required for XLSX export. Install it with: pip install openpyxl")
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Registers"
        
        # Get all unique column names from all registers
        column_names = set()
        for reg in registers:
            column_names.update(reg.keys())
        column_names = sorted(list(column_names))
        
        # Write headers
        for col_idx, col_name in enumerate(column_names, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.value = col_name
            # Bold header
            cell.font = Font(bold=True)
        
        # Write data rows
        for row_idx, reg in enumerate(registers, 2):
            for col_idx, col_name in enumerate(column_names, 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.value = reg.get(col_name, "")
        
        # Auto-adjust column widths
        for col_idx, col_name in enumerate(column_names, 1):
            max_length = len(str(col_name)) + 2
            for row_idx in range(2, len(registers) + 2):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)) + 2)
            ws.column_dimensions[ws.cell(1, col_idx).column_letter].width = min(max_length, 50)
        
        wb.save(filename)
        print(f"✓ Exported {len(registers)} registers to {filename}")
        return filename


class RegisterImporter:
    """Import register tables from JSON and XLSX formats"""
    
    @staticmethod
    def import_from_json(filename):
        """Import registers from JSON file"""
        with open(filename, 'r', encoding='utf-8') as f:
            data = json.load(f)
        
        # Ensure readOnStart column exists
        for reg in data:
            if 'readOnStart' not in reg:
                reg['readOnStart'] = False
        
        print(f"✓ Imported {len(data)} registers from {filename}")
        return tuple(data)
    
    @staticmethod
    def import_from_xlsx(filename):
        """Import registers from XLSX file"""
        if not XLSX_AVAILABLE:
            raise ImportError("openpyxl is required for XLSX import. Install it with: pip install openpyxl")
        
        wb = load_workbook(filename)
        ws = wb.active
        
        # Read headers
        headers = []
        for cell in ws[1]:
            headers.append(cell.value)
        
        # Read data rows
        data = []
        for row in ws.iter_rows(min_row=2, values_only=False):
            row_data = {}
            for col_idx, header in enumerate(headers):
                cell = row[col_idx]
                value = cell.value
                # Convert boolean strings
                if value == "TRUE" or value is True:
                    value = True
                elif value == "FALSE" or value is False:
                    value = False
                # Try to convert numbers
                elif isinstance(value, str):
                    try:
                        value = int(value)
                    except ValueError:
                        try:
                            value = float(value)
                        except ValueError:
                            pass
                
                if header and value is not None:
                    row_data[header] = value
            
            if row_data:  # Only add non-empty rows
                data.append(row_data)
        
        # Ensure readOnStart column exists
        for reg in data:
            if 'readOnStart' not in reg:
                reg['readOnStart'] = False
        
        print(f"✓ Imported {len(data)} registers from {filename}")
        return tuple(data)


def export_all_registers(output_dir="registers_export", format_type="both"):
    """
    Export all register tables to files
    
    Args:
        output_dir: Directory to save export files
        format_type: 'json', 'xlsx', or 'both'
    """
    Path(output_dir).mkdir(exist_ok=True)
    
    tables = {
        "datalogger": datalogger_registers_table,
        "inverter_timedate": inverter_timedate_registers_table,
        "inverter": inverter_registers_table
    }
    
    exporter = RegisterExporter()
    
    for table_name, table_data in tables.items():
        if format_type in ("json", "both"):
            json_file = os.path.join(output_dir, f"{table_name}.json")
            exporter.export_to_json(table_data, json_file)
        
        if format_type in ("xlsx", "both"):
            if XLSX_AVAILABLE:
                xlsx_file = os.path.join(output_dir, f"{table_name}.xlsx")
                exporter.export_to_xlsx(table_data, xlsx_file)
            else:
                print(f"⚠ Skipping XLSX export - openpyxl not installed")


def import_all_registers(input_dir="registers", format_type="json"):
    """
    Import all register tables from files by walking the directory
    
    Args:
        input_dir: Directory containing import files
        format_type: 'json' or 'xlsx'
    """
    importer = RegisterImporter()
    tables = {}
    
    if not os.path.exists(input_dir):
        print(f"⚠ Directory not found: {input_dir}")
        return tables
    
    # Walk through directory and find matching files
    for root, dirs, files in os.walk(input_dir):
        for filename in files:
            if format_type == "json" and filename.endswith(".json"):
                file_path = os.path.join(root, filename)
                table_name = os.path.splitext(filename)[0]
                tables[table_name] = importer.import_from_json(file_path)
            elif format_type == "xlsx" and filename.endswith(".xlsx"):
                file_path = os.path.join(root, filename)
                table_name = os.path.splitext(filename)[0]
                tables[table_name] = importer.import_from_xlsx(file_path)
    
    if not tables:
        print(f"⚠ No {format_type.upper()} files found in {input_dir}")
    
    return tables


def convert_registers(input_dir, input_format, output_dir, output_format):
    """
    Convert register tables between JSON and XLSX formats
    
    Args:
        input_dir: Directory containing source files
        input_format: 'json' or 'xlsx'
        output_dir: Directory to save converted files
        output_format: 'json' or 'xlsx'
    """
    if input_format == output_format:
        print(f"⚠ Input and output formats are the same ({input_format})")
        return
    
    # Import from source format
    print(f"Reading {input_format.upper()} files from '{input_dir}'...")
    tables = import_all_registers(input_dir, input_format)
    
    if not tables:
        print(f"✗ No registers imported from '{input_dir}'")
        return
    
    # Export to target format
    print(f"Converting to {output_format.upper()} and saving to '{output_dir}'...")
    Path(output_dir).mkdir(exist_ok=True)
    exporter = RegisterExporter()
    
    for table_name, table_data in tables.items():
        if output_format == "json":
            json_file = os.path.join(output_dir, f"{table_name}.json")
            exporter.export_to_json(table_data, json_file)
        else:
            if XLSX_AVAILABLE:
                xlsx_file = os.path.join(output_dir, f"{table_name}.xlsx")
                exporter.export_to_xlsx(table_data, xlsx_file)
            else:
                print(f"✗ openpyxl not installed - cannot export to XLSX")
                return
    
    print(f"\n✓ Successfully converted {len(tables)} register tables from {input_format.upper()} to {output_format.upper()}")


if __name__ == "__main__":
    import sys
    
    if len(sys.argv) < 2:
        print("Register Tool - Export/Import/Convert utility")
        print("\nUsage:")
        print("  python registerTool.py export [json|xlsx|both] [output_dir]")
        print("  python registerTool.py import [json|xlsx] [input_dir]")
        print("  python registerTool.py convert <input_format> <output_format> [input_dir] [output_dir]")
        print("\nExamples:")
        print("  python registerTool.py export both ./exports")
        print("  python registerTool.py import json ./exports")
        print("  python registerTool.py convert json xlsx ./json_files ./xlsx_files")
        print("  python registerTool.py convert xlsx json ./xlsx_files ./json_files")
        sys.exit(1)
    
    command = sys.argv[1]
    
    if command == "export":
        fmt = sys.argv[2] if len(sys.argv) > 2 else "both"
        out_dir = sys.argv[3] if len(sys.argv) > 3 else "registers_export"
        export_all_registers(out_dir, fmt)
        print(f"\n✓ All registers exported to '{out_dir}' in {fmt} format")
    
    elif command == "import":
        fmt = sys.argv[2] if len(sys.argv) > 2 else "json"
        in_dir = sys.argv[3] if len(sys.argv) > 3 else "registers_export"
        tables = import_all_registers(in_dir, fmt)
        print(f"\n✓ All registers imported from '{in_dir}' ({fmt} format)")
        for name, data in tables.items():
            print(f"  - {name}: {len(data)} registers")
    
    elif command == "convert":
        if len(sys.argv) < 4:
            print("✗ Convert requires input and output formats")
            print("Usage: python registerTool.py convert <input_format> <output_format> [input_dir] [output_dir]")
            sys.exit(1)
        
        input_fmt = sys.argv[2].lower()
        output_fmt = sys.argv[3].lower()
        input_dir = sys.argv[4] if len(sys.argv) > 4 else "registers_export"
        output_dir = sys.argv[5] if len(sys.argv) > 5 else f"registers_convert_{output_fmt}"
        
        if input_fmt not in ("json", "xlsx") or output_fmt not in ("json", "xlsx"):
            print("✗ Format must be 'json' or 'xlsx'")
            sys.exit(1)
        
        convert_registers(input_dir, input_fmt, output_dir, output_fmt)
    
    else:
        print(f"Unknown command: {command}")
        print("\nAvailable commands: export, import, convert")
        sys.exit(1) 

